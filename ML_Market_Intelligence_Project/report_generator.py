import io
from fpdf import FPDF
from docx import Document

def generate_pdf_report(project_data, risk_data, swot_data, agent_result, market_data):
    pdf = FPDF()
    pdf.add_page()
    
    # Title
    pdf.set_font("helvetica", 'B', 16)
    pdf.cell(0, 10, "Failure Prediction AI - Comprehensive Assessment", align='C')
    pdf.ln(15)
    
    # Project Info
    pdf.set_font("helvetica", 'B', 14)
    pdf.cell(0, 10, f"Startup: {project_data.get('startup_name', 'N/A')}")
    pdf.ln(8)
    pdf.set_font("helvetica", '', 12)
    pdf.cell(0, 10, f"Industry: {project_data.get('industry', 'N/A')} | Budget: ${project_data.get('budget', 0):,}")
    pdf.ln(10)
    pdf.multi_cell(0, 8, f"Description: {project_data.get('description', 'N/A')}")
    pdf.ln(10)
    
    # Risk Assessment
    pdf.set_font("helvetica", 'B', 14)
    pdf.cell(0, 10, "Risk Assessment")
    pdf.ln(8)
    pdf.set_font("helvetica", '', 12)
    for k, v in risk_data.items():
        pdf.cell(0, 8, f"- {k}: {v}/5")
        pdf.ln(6)
    pdf.ln(5)
    
    # Market Analysis
    pdf.set_font("helvetica", 'B', 14)
    pdf.cell(0, 10, "Market Analysis")
    pdf.ln(8)
    pdf.set_font("helvetica", '', 12)
    pdf.cell(0, 8, f"- Total Addressable Market (TAM): ${market_data.get('tam', 0)}B")
    pdf.ln(6)
    pdf.cell(0, 8, f"- Serviceable Addressable Market (SAM): ${market_data.get('sam', 0)}M")
    pdf.ln(6)
    pdf.cell(0, 8, f"- Serviceable Obtainable Market (SOM): ${market_data.get('som', 0)}M")
    pdf.ln(10)
    
    # AI Recommendations
    if agent_result:
        pdf.set_font("helvetica", 'B', 14)
        pdf.cell(0, 10, "Strategic AI Recommendations")
        pdf.ln(8)
        pdf.set_font("helvetica", '', 11)
        for rec in agent_result.get("recommendations", [])[:4]:
            pdf.multi_cell(0, 8, f"• {rec.get('title')} ({rec.get('priority')}): {rec.get('description')}")
            pdf.ln(4)
            
    return bytes(pdf.output())

def generate_docx_report(project_data, risk_data, swot_data, agent_result, market_data):
    doc = Document()
    doc.add_heading('Failure Prediction AI - Comprehensive Assessment', 0)
    
    doc.add_heading(f"Startup: {project_data.get('startup_name', 'N/A')}", level=1)
    doc.add_paragraph(f"Industry: {project_data.get('industry', 'N/A')} | Budget: ${project_data.get('budget', 0):,}")
    doc.add_paragraph(f"Description: {project_data.get('description', 'N/A')}")
    
    doc.add_heading('Risk Assessment', level=2)
    for k, v in risk_data.items():
        doc.add_paragraph(f"{k}: {v}/5", style='List Bullet')
        
    doc.add_heading('Market Analysis', level=2)
    doc.add_paragraph(f"TAM: ${market_data.get('tam', 0)}B", style='List Bullet')
    doc.add_paragraph(f"SAM: ${market_data.get('sam', 0)}M", style='List Bullet')
    doc.add_paragraph(f"SOM: ${market_data.get('som', 0)}M", style='List Bullet')
    
    if agent_result:
        doc.add_heading('Strategic AI Recommendations', level=2)
        for rec in agent_result.get("recommendations", [])[:4]:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{rec.get('title')} ({rec.get('priority')}): ").bold = True
            p.add_run(f"{rec.get('description')}")
            
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
